"""The first-run inference API: GET /infer, POST /infer/ai-fill, POST /infer/apply, status.setup."""

import json
from collections.abc import Iterator
from pathlib import Path
from typing import Any

import keyring
import pytest
from docx import Document
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
from fastapi.testclient import TestClient
from test_vault import CHEAP, MemoryKeyring

from linkit_all.api.app import create_app
from linkit_all.apply.vault import Vault
from linkit_all.core.config import Settings
from linkit_all.providers import ProviderError
from linkit_all.providers.base import LLMProvider

BASE = "http://127.0.0.1:8765"
H = {"origin": BASE}


@pytest.fixture(autouse=True)
def mem_keyring() -> Iterator[MemoryKeyring]:
    previous = keyring.get_keyring()
    backend = MemoryKeyring()
    keyring.set_keyring(backend)
    yield backend
    keyring.set_keyring(previous)


@pytest.fixture
def client(tmp_path: Path) -> Iterator[TestClient]:
    app = create_app(Settings(data_dir=tmp_path / "data"))
    app.state.vault = Vault(app.state.db, "api-test", CHEAP)
    with TestClient(app, base_url=BASE) as c:
        yield c


def cv_bytes(tmp_path: Path) -> bytes:
    doc = Document()
    for line in [
        "Sara Ahmed",
        "Cairo, Egypt | +20 100 123 4567 | sara.ahmed@example.com",
        "Portfolio: behance.net/sara",
        "Profiles: ",
        "EXPERIENCE",
        "Backend Engineer at Acme Corp | Jan 2021 - Present",
        "Software Developer - Globex | 03/2018 - 12/2020",
        "SKILLS",
        "Python, FastAPI, PostgreSQL, Docker",
    ]:
        doc.add_paragraph(line)
    rid = doc.part.relate_to(
        "https://www.linkedin.com/in/sara-ahmed", RT.HYPERLINK, is_external=True
    )
    link = parse_xml(
        f'<w:hyperlink {nsdecls("w", "r")} r:id="{rid}"><w:r><w:t>LinkedIn</w:t></w:r></w:hyperlink>'
    )
    doc.paragraphs[3]._p.append(link)
    path = tmp_path / "sara.docx"
    doc.save(str(path))
    return path.read_bytes()


def upload(client: TestClient, tmp_path: Path) -> None:
    r = client.post(
        "/api/profile/cv",
        params={"filename": "sara.docx"},
        content=cv_bytes(tmp_path),
        headers=H,
    )
    assert r.status_code == 200, r.text


def connect_ai(client: TestClient) -> None:
    body = {"ai": {"provider": "claude_cli"}}
    assert client.put("/api/search", json=body, headers=H).status_code == 200


def apply(client: TestClient, **body: Any) -> Any:
    return client.post("/api/infer/apply", json=body, headers=H)


def test_get_infer_before_and_after_upload(client: TestClient, tmp_path: Path) -> None:
    empty = client.get("/api/infer").json()
    assert empty["values"] == {} and empty["has_cv"] is False
    assert empty["country"]["name"] == ""
    upload(client, tmp_path)
    got = client.get("/api/infer").json()
    v = got["values"]
    assert got["has_cv"] and v["profile.full_name"]["value"] == "Sara Ahmed"
    assert v["profile.email"]["value"] == "sara.ahmed@example.com"
    assert (
        v["answers.linkedin"]["value"] == "https://www.linkedin.com/in/sara-ahmed"
    )  # from the link
    assert v["answers.country"]["value"] == "Egypt"
    assert v["track"]["value"] == "tech" and v["track"]["level"] == "suggest"
    assert v["answers.authorized_countries"]["level"] == "suggest"
    assert got["country"] == {
        "name": "Egypt",
        "jobspy": "egypt",
        "currency": "EGP",
        "region": "egypt",
    }
    assert "wuzzuf" in got["sources"]["regional"] and "remoteok" in got["sources"]["enabled"]


def test_get_infer_is_read_only(client: TestClient, tmp_path: Path) -> None:
    upload(client, tmp_path)
    before = (client.get("/api/profile").json(), client.get("/api/search").json())
    before_answers = client.get("/api/profile/answers").json()
    assert client.get("/api/infer").status_code == 200
    assert (client.get("/api/profile").json(), client.get("/api/search").json()) == before
    assert client.get("/api/profile/answers").json() == before_answers


def test_apply_writes_only_given_keys(client: TestClient, tmp_path: Path) -> None:
    upload(client, tmp_path)
    keep = {
        "full_name": "Keep Me",
        "email": "keep@x.io",
        "skills": ["Go"],
        "locked_facts": ["Acme"],
    }
    client.put("/api/profile", json=keep, headers=H)
    r = apply(
        client,
        values={
            "profile.phone": "+201001234567",
            "answers.city": "Cairo",
            "search.queries": ["Backend Engineer"],
            "filters.seniority": ["senior", "lead"],
        },
    )
    assert r.status_code == 200, r.text
    profile = client.get("/api/profile").json()
    assert profile["phone"] == "+201001234567"
    assert {k: profile[k] for k in keep} == keep  # untouched
    answers = client.get("/api/profile/answers").json()
    assert answers["city"] == "Cairo" and answers["first_name"] == ""
    assert answers["authorized_countries"] == []  # never written unless given
    cfg = client.get("/api/search").json()
    assert cfg["search"]["queries"] == ["Backend Engineer"]
    assert cfg["filters"]["seniority"] == ["senior", "lead"]
    assert cfg["filters"]["title_exclude"] == [] and cfg["sources"]["jobspy"] == []
    assert r.json()["setup"]["roles"] is True


def test_authorized_countries_only_when_given(client: TestClient, tmp_path: Path) -> None:
    upload(client, tmp_path)
    assert apply(client, values={"answers.country": "Egypt"}).status_code == 200
    assert client.get("/api/profile/answers").json()["authorized_countries"] == []
    assert apply(client, values={"answers.authorized_countries": ["Egypt"]}).status_code == 200
    assert client.get("/api/profile/answers").json()["authorized_countries"] == ["Egypt"]


def test_apply_without_profile_row_creates_it(client: TestClient) -> None:
    assert apply(client, values={"answers.city": "Cairo"}).status_code == 200
    assert client.get("/api/profile/answers").json()["city"] == "Cairo"


def test_where_remote_and_track_presets(client: TestClient) -> None:
    mine = {"sources": {"enabled": ["my_board"], "jobspy": ["indeed"]}}
    client.put("/api/search", json=mine, headers=H)
    r = apply(client, where=["Remote Hub", "Cairo, Egypt"], remote=True, track="tech")
    assert r.status_code == 200, r.text
    cfg = client.get("/api/search").json()
    assert cfg["search"]["locations"] == ["Remote Hub", "Cairo, Egypt", "Remote"]
    assert cfg["filters"]["locations"] == ["remote hub", "cairo, egypt", "remote"]
    assert cfg["filters"]["work_modes"] == ["remote", "hybrid", "onsite"]
    assert cfg["filters"]["salary_currency"] == "EGP"
    assert cfg["sources"]["jobspy_country"] == "egypt"
    assert cfg["sources"]["jobspy"] == ["linkedin", "indeed", "google"]
    assert "wuzzuf" in cfg["sources"]["regional"]
    assert cfg["sources"]["enabled"][0] == "my_board"  # union: nothing of the user's is removed
    assert "remoteok" in cfg["sources"]["enabled"]
    assert cfg["sources"]["telegram_channels"] == []  # still the default list: replaced
    assert client.get("/api/profile/answers").json()["salary_currency"] == "EGP"
    assert r.json()["setup"]["where"] is True
    # not remote -> hybrid/onsite only; an edited channel list and a set currency stay
    client.put(
        "/api/search",
        json={**cfg, "sources": {**cfg["sources"], "telegram_channels": ["mychannel"]}},
        headers=H,
    )
    client.put("/api/profile/answers", json={"salary_currency": "USD"}, headers=H)
    assert apply(client, where=["Dubai, UAE"], remote=False, track="general").status_code == 200
    cfg = client.get("/api/search").json()
    assert cfg["filters"]["work_modes"] == ["hybrid", "onsite"]
    assert cfg["search"]["locations"] == ["Dubai, UAE"]
    assert cfg["sources"]["telegram_channels"] == ["mychannel"]
    assert cfg["filters"]["salary_currency"] == "AED"
    assert client.get("/api/profile/answers").json()["salary_currency"] == "USD"


def test_apply_rejects_unknown_and_invalid(client: TestClient) -> None:
    r = apply(client, values={"profile.nope": "x"})
    assert r.status_code == 422 and "profile.nope" in r.json()["detail"]
    assert apply(client, values={"track": "tech"}).status_code == 422
    assert apply(client, values={"filters.seniority": ["wizard"]}).status_code == 422
    assert apply(client, values={"profile.skills": "not a list"}).status_code == 422
    assert apply(client, track="nonsense").status_code == 422
    assert client.get("/api/search").json()["filters"]["seniority"] == []  # nothing half-written


class FakeLLM(LLMProvider):
    name = "fake"

    def __init__(self, reply: str | Exception) -> None:
        super().__init__("fake-model")
        self.reply = reply
        self.prompts: list[str] = []

    async def _call(self, prompt: str) -> str:
        self.prompts.append(prompt)
        if isinstance(self.reply, Exception):
            raise self.reply
        return self.reply


def patch_provider(monkeypatch: pytest.MonkeyPatch, fake: FakeLLM) -> None:
    monkeypatch.setattr("linkit_all.providers.get_provider", lambda *a, **k: fake)


def test_ai_fill_keeps_only_values_found_in_the_cv(
    client: TestClient, tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    upload(client, tmp_path)
    connect_ai(client)
    reply = {
        "portfolio": "behance.net/sara",  # in the CV
        "github": "github.com/saraahmed",  # not in the CV: dropped
        "full_name": "Someone Else",  # the rules already found a name: not asked, ignored
        "current_title": "Chief Wizard",  # not in the CV: dropped
    }
    fake = FakeLLM(json.dumps(reply))
    patch_provider(monkeypatch, fake)
    before = client.get("/api/infer").json()["values"]
    assert "answers.portfolio" not in before
    r = client.post("/api/infer/ai-fill", headers=H)
    assert r.status_code == 200, r.text
    values = r.json()["values"]
    assert values["answers.portfolio"] == {
        "value": "behance.net/sara", "source": "ai", "level": "suggest", "check": True,
    }  # fmt: skip
    assert "answers.github" not in values
    assert values["answers.current_title"]["source"] == "cv"  # found by the rules, not asked
    assert values["profile.full_name"]["value"] == "Sara Ahmed"  # rules win
    assert len(fake.prompts) == 1
    assert "Sara Ahmed" in fake.prompts[0]
    for sensitive in ("sponsor", "salary", "visa", "gender", "disability"):
        assert sensitive not in fake.prompts[0].lower()
    # saved next to the CV: GET returns it without another call
    assert client.get("/api/infer").json()["values"]["answers.portfolio"]["source"] == "ai"
    assert len(fake.prompts) == 1
    # a new CV invalidates it
    other = Document()
    other.add_paragraph("Omar Khaled")
    other.add_paragraph("omar@example.com")
    path = tmp_path / "omar.docx"
    other.save(str(path))
    client.post(
        "/api/profile/cv", params={"filename": "omar.docx"}, content=path.read_bytes(), headers=H
    )
    assert "answers.portfolio" not in client.get("/api/infer").json()["values"]


def test_ai_fill_skills_must_each_appear_in_the_cv(
    client: TestClient, monkeypatch: pytest.MonkeyPatch
) -> None:
    doc = Document()
    for line in ["Omar Khaled", "omar@example.com", "Worked with Kubernetes and Terraform"]:
        doc.add_paragraph(line)
    path = client.app.state.settings.data_dir / "omar.docx"  # type: ignore[attr-defined]
    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(path))
    client.post(
        "/api/profile/cv", params={"filename": "omar.docx"}, content=path.read_bytes(), headers=H
    )
    connect_ai(client)
    patch_provider(monkeypatch, FakeLLM(json.dumps({"skills": ["kubernetes", "Quantum Magic"]})))
    r = client.post("/api/infer/ai-fill", headers=H)
    skills = r.json()["values"]["profile.skills"]
    assert skills["source"] in ("ai", "cv")  # the lexicon may already have found some
    assert "Quantum Magic" not in skills["value"]


def test_ai_fill_errors(
    client: TestClient, tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    assert client.post("/api/infer/ai-fill", headers=H).status_code == 404  # no CV
    upload(client, tmp_path)
    r = client.post("/api/infer/ai-fill", headers=H)  # no provider chosen
    assert r.status_code == 409 and "AI" in r.json()["detail"]
    connect_ai(client)
    patch_provider(monkeypatch, FakeLLM(ProviderError("sign in to Claude first")))
    r = client.post("/api/infer/ai-fill", headers=H)
    assert r.status_code == 502 and "sign in" in r.json()["detail"]


def test_status_setup_and_next_step(client: TestClient, tmp_path: Path) -> None:
    def setup() -> dict[str, Any]:
        return client.get("/api/status").json()["setup"]  # type: ignore[no-any-return]

    assert setup() == {
        "cv": False, "ai": False, "profile": False, "roles": False, "where": False,
        "required_done": False, "next_step": "cv",
    }  # fmt: skip
    upload(client, tmp_path)
    assert setup()["cv"] and setup()["next_step"] == "ai"
    connect_ai(client)
    assert setup()["ai"] and setup()["next_step"] == "profile"
    apply(client, values={"profile.full_name": "Sara Ahmed", "profile.email": "s@x.io"})
    assert setup()["profile"] and setup()["next_step"] == "roles"
    apply(client, values={"search.queries": ["Backend Engineer"]})
    assert setup()["roles"] and setup()["next_step"] == "where"
    status = apply(client, where=["Cairo, Egypt"]).json()
    assert status["setup"]["required_done"] is True and status["setup"]["next_step"] is None
    assert (
        status["profile_ready"] and status["cv_ready"] and status["search_ready"]
    )  # unchanged keys
