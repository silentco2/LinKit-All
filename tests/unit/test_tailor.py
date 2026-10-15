import asyncio
import json
from pathlib import Path

import docx
import pymupdf
import pytest
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

from linkit_all.core.config import Settings
from linkit_all.core.db import init_db
from linkit_all.core.models import Job, JobEvent, JobState
from linkit_all.core.profile import ProfileIn, import_profile
from linkit_all.providers.base import LLMProvider
from linkit_all.tailor.docx.segments import PatchError, apply_patch, load_segments
from linkit_all.tailor.patches import Patch, Rules, plan_patches, vocabulary
from linkit_all.tailor.pipeline import tailor_docx, tailor_job

NS = (
    'xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006" '
    'xmlns:wps="http://schemas.microsoft.com/office/word/2010/wordprocessingShape" '
    'xmlns:v="urn:schemas-microsoft-com:vml"'
)
TEXTBOX = (
    f"<w:r {nsdecls('w')} {NS}><mc:AlternateContent>"
    "<mc:Choice Requires='wps'><w:drawing><wps:txbx><w:txbxContent>"
    "<w:p><w:r><w:t>Skills: Python</w:t></w:r></w:p>"
    "</w:txbxContent></wps:txbx></w:drawing></mc:Choice>"
    "<mc:Fallback><w:pict><v:textbox><w:txbxContent>"
    "<w:p><w:r><w:t>Skills: Python</w:t></w:r></w:p>"
    "</w:txbxContent></v:textbox></w:pict></mc:Fallback>"
    "</mc:AlternateContent></w:r>"
)


def make_cv(path: Path) -> Path:
    d = docx.Document()
    d.sections[0].header.paragraphs[0].text = "Jane Doe - jane@example.com"
    d.add_paragraph("Software Engineer at Acme, 2019-2023")
    p = d.add_paragraph()
    p.add_run("Built scal").bold = True
    p.add_run("able backend services in Python and Go")
    d.add_paragraph("Improved latency by 40% for 2M users")
    d.add_table(rows=1, cols=1).cell(0, 0).text = "Mentored junior developers"
    d.add_paragraph()._p.append(parse_xml(TEXTBOX))
    d.save(str(path))
    return path


def texts(path: Path) -> dict[int, str]:
    return {s.id: s.text for s in load_segments(docx.Document(str(path)))}


def seg_id(path: Path, startswith: str) -> int:
    return next(i for i, t in texts(path).items() if t.startswith(startswith))


def fake_render(src: Path, out_dir: Path) -> Path:
    """Stand-in for LibreOffice: one PDF page with each segment's text."""
    out = pymupdf.open()
    page = out.new_page()
    for i, t in enumerate(t for t in texts(src).values() if t):
        page.insert_text((50, 60 + 16 * i), t, fontsize=10)
    pdf = out_dir / f"{src.stem}.pdf"
    out.save(str(pdf))
    return pdf


# --- segments -----------------------------------------------------------------


def test_segments_cover_header_table_and_textbox_once(tmp_path: Path) -> None:
    all_text = list(texts(make_cv(tmp_path / "cv.docx")).values())
    assert "Jane Doe - jane@example.com" in all_text
    assert "Mentored junior developers" in all_text
    assert all_text.count("Skills: Python") == 1  # mc:Fallback copy is not a second segment


def test_apply_patch_keeps_first_run_format_and_updates_textbox_fallback(tmp_path: Path) -> None:
    d = docx.Document(str(make_cv(tmp_path / "cv.docx")))
    segs = {s.text: s for s in load_segments(d)}
    apply_patch(segs["Built scalable backend services in Python and Go"], "scalable", "reliable")
    runs = [(r.text, r.bold) for r in d.paragraphs[1].runs]
    assert runs == [("Built reliable", True), (" backend services in Python and Go", None)]
    box = segs["Skills: Python"]
    apply_patch(box, "Skills", "Tools")
    assert len(box.paragraphs) == 2
    assert all("Tools: Python" in "".join(p.itertext()) for p in box.paragraphs)


def test_apply_patch_rejects_ambiguous_or_missing_text(tmp_path: Path) -> None:
    d = docx.Document(str(make_cv(tmp_path / "cv.docx")))
    seg = next(s for s in load_segments(d) if s.text.startswith("Built"))
    with pytest.raises(PatchError):
        apply_patch(seg, "not there", "x")
    with pytest.raises(PatchError):
        apply_patch(seg, "e", "x")  # appears many times
    assert seg.text.startswith("Built scalable")


# --- rules --------------------------------------------------------------------

SEGS = {0: "Software Engineer at Acme, 2019-2023", 1: "Built scalable backend services in Python"}


def rules(**kw: object) -> Rules:
    vocab = vocabulary(" ".join(SEGS.values()), ["Docker"])
    return Rules(**{"locked_facts": ["Acme"], "allowed_terms": vocab, **kw})  # type: ignore[arg-type]


@pytest.mark.parametrize(
    ("patch", "reason"),
    [
        (Patch(segment_id=0, old_text="Acme", new_text="Globex"), "locked fact"),
        (Patch(segment_id=0, old_text="2019", new_text="2018"), "number"),
        (Patch(segment_id=1, old_text="Python", new_text="Kubernetes"), "Kubernetes"),
        (Patch(segment_id=1, old_text="Built", new_text="Kubernetes"), "Kubernetes"),
        (Patch(segment_id=1, old_text="backend", new_text="backend and AWS"), "AWS"),
        (Patch(segment_id=1, old_text="scalable", new_text="highly scalable, resilient"), "grows"),
        (Patch(segment_id=1, old_text="e", new_text="x"), "exactly once"),
        (Patch(segment_id=9, old_text="a", new_text="b"), "unknown segment"),
        (Patch(segment_id=1, old_text="Built", new_text="Built\tit"), "tab"),
    ],
)
def test_unsafe_patches_rejected(patch: Patch, reason: str) -> None:
    plan = plan_patches(SEGS, [patch], rules())
    assert not plan.accepted
    assert reason in plan.rejected[0][1]


def test_safe_rephrasing_and_listed_skill_accepted() -> None:
    patches = [
        Patch(segment_id=1, old_text="Built", new_text="Designed"),
        Patch(segment_id=1, old_text="in Python", new_text="in Docker"),
    ]
    plan = plan_patches(SEGS, patches, rules(max_growth=0.2, max_word_change=1))
    assert plan.accepted == patches


def test_word_change_cap() -> None:
    patches = [
        Patch(segment_id=1, old_text="scalable", new_text="reliable"),
        Patch(segment_id=1, old_text="backend", new_text="server"),
    ]
    plan = plan_patches(SEGS, patches, rules(max_word_change=0.1))  # 11 words -> 1 change
    assert plan.accepted == patches[:1]
    assert "cap" in plan.rejected[0][1]


# --- pipeline -----------------------------------------------------------------


def test_tailor_docx_applies_and_verifies(tmp_path: Path) -> None:
    cv = make_cv(tmp_path / "cv.docx")
    base_pdf = fake_render(cv, tmp_path)
    sid = seg_id(cv, "Built")
    patch = Patch(segment_id=sid, old_text="scalable", new_text="reliable")
    bad = Patch(segment_id=sid, old_text="Go", new_text="Rust")
    rules_ = Rules(["Acme"], vocabulary(" ".join(texts(cv).values()), []))
    out = tailor_docx(
        cv, base_pdf, [patch, bad], rules_, tmp_path / "out", min_ssim=0.5, render=fake_render
    )
    assert out.applied == [patch]
    assert out.rejected[0][0] == bad
    assert out.verify is not None and out.verify.ok
    assert "reliable backend" in texts(out.docx)[sid]
    assert texts(cv)[sid].startswith("Built scalable")  # base CV untouched


def test_tailor_docx_rolls_back_when_layout_breaks(tmp_path: Path) -> None:
    cv = make_cv(tmp_path / "cv.docx")
    base_pdf = fake_render(cv, tmp_path)
    patch = Patch(segment_id=seg_id(cv, "Built"), old_text="scalable", new_text="reliable")
    rules_ = Rules([], vocabulary(" ".join(texts(cv).values()), []))
    out = tailor_docx(
        cv, base_pdf, [patch], rules_, tmp_path / "out", min_ssim=1.01, render=fake_render
    )
    assert out.applied == []
    assert "rolled back" in out.rejected[0][1]
    assert texts(out.docx) == texts(cv)
    assert out.pdf.read_bytes() == base_pdf.read_bytes()


class FakeProvider(LLMProvider):
    name = "fake"

    def __init__(self, reply: str) -> None:
        super().__init__("fake")
        self.reply = reply

    async def _call(self, prompt: str) -> str:
        assert 'Never change these locked facts: "Acme"' in prompt
        return self.reply


def test_tailor_job_stores_version_and_transitions(tmp_path: Path) -> None:
    settings = Settings(data_dir=tmp_path / "data")
    db = init_db(settings.db_path)
    cv = make_cv(tmp_path / "cv.docx")
    import_profile(
        db,
        ProfileIn(full_name="Jane", email="j@x.io", cv_path=cv, locked_facts=["Acme", " "]),
        settings.data_dir,
    )
    with db.begin() as s:
        job = Job(source="t", url="u", title="Backend", company="Initech", dedupe_hash="h")
        job.state = JobState.queued_apply
        s.add(job)
    sid = seg_id(cv, "Built")
    patch = {"segment_id": sid, "old_text": "scalable", "new_text": "reliable"}
    reply = json.dumps({"patches": [patch]})

    row = asyncio.run(tailor_job(db, job.id, FakeProvider(reply), settings, render=fake_render))
    assert [p["new_text"] for p in row.applied] == ["reliable"]
    assert Path(row.pdf_path).exists()
    with db() as s:
        assert s.get_one(Job, job.id).state == JobState.tailored
        event = s.query(JobEvent).filter_by(job_id=job.id).one()
        assert event.artifacts == [row.docx_path, row.pdf_path]
    again = asyncio.run(tailor_job(db, job.id, FakeProvider("{}"), settings, render=fake_render))
    assert again.id == row.id  # idempotent: no second AI call result stored
