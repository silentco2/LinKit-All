"""Tailor stage: AI proposes patches, code validates, applies, renders and verifies."""

import asyncio
import shutil
from collections.abc import Callable
from dataclasses import dataclass
from pathlib import Path
from typing import Protocol

from docx import Document
from sqlalchemy import select
from sqlalchemy.orm import Session, sessionmaker

from linkit_all.core.config import Settings
from linkit_all.core.models import Job, JobState, LockedFact, Profile, TailoredCV
from linkit_all.core.state import transition
from linkit_all.providers.base import LLMProvider
from linkit_all.providers.base import render as render_prompt
from linkit_all.tailor import latex
from linkit_all.tailor.docx.segments import PatchError, apply_patch, load_segments
from linkit_all.tailor.patches import Patch, PatchList, Rules, plan_patches, vocabulary
from linkit_all.tailor.render.latex import render_pdf_latex
from linkit_all.tailor.render.libreoffice import render_pdf
from linkit_all.tailor.verify.layout import VerifyReport, verify_pdf

Renderer = Callable[[Path, Path], Path]


class TailorError(RuntimeError):
    pass


@dataclass
class TailorResult:
    docx: Path
    pdf: Path
    applied: list[Patch]
    rejected: list[tuple[Patch, str]]
    verify: VerifyReport | None


class CvSegment(Protocol):
    @property
    def id(self) -> int: ...
    @property
    def text(self) -> str: ...


def cv_segments(cv: Path) -> list[CvSegment]:
    """Non-empty text segments of the base CV (DOCX or LaTeX source)."""
    if cv.suffix.lower() == ".tex":
        return list(latex.load_segments(cv.read_bytes().decode("utf-8")))
    return [s for s in load_segments(Document(str(cv))) if s.text.strip()]


def renderer_for(cv: Path) -> Renderer:
    return render_pdf_latex if cv.suffix.lower() == ".tex" else render_pdf


def ensure_base_pdf(cv: Path, render: Renderer | None = None) -> Path:
    """Render the untouched CV once; it is the layout reference and the fallback output."""
    pdf = cv.with_suffix(".pdf")
    if not pdf.exists() or pdf.stat().st_mtime < cv.stat().st_mtime:
        pdf = (render or renderer_for(cv))(cv, cv.parent)
    return pdf


# Writes the patched CV to `dest` when at least one patch applies; returns (applied, failed).
Builder = Callable[[list[Patch], Path], tuple[list[Patch], list[tuple[Patch, str]]]]


def _tailor(
    base: Path,
    base_pdf: Path,
    texts: dict[int, str],
    patches: list[Patch],
    rules: Rules,
    out_dir: Path,
    build: Builder,
    min_ssim: float,
    render: Renderer,
) -> TailorResult:
    out_dir.mkdir(parents=True, exist_ok=True)
    plan = plan_patches(texts, patches, rules)

    # ponytail: one fallback attempt with only non-growing patches (growth is what breaks
    # layout); bisect patches if real CVs show that is too coarse.
    attempts = [plan.accepted]
    shrinking = [p for p in plan.accepted if len(p.new_text) <= len(p.old_text)]
    if shrinking and len(shrinking) < len(plan.accepted):
        attempts.append(shrinking)

    report: VerifyReport | None = None
    for attempt in attempts:
        dest = out_dir / f"cv{base.suffix}"
        applied, failed = build(attempt, dest)
        if not applied:
            plan.rejected += failed
            break
        pdf = render(dest, out_dir)
        report = verify_pdf(base_pdf, pdf, min_ssim)
        if report.ok:
            dropped = [(p, "dropped to keep the layout") for p in plan.accepted if p not in attempt]
            return TailorResult(dest, pdf, applied, plan.rejected + failed + dropped, report)

    # Nothing safe to apply, or layout broke: submit the untouched CV.
    shutil.copyfile(base, out_dir / f"cv{base.suffix}")
    shutil.copyfile(base_pdf, out_dir / "cv.pdf")
    why = f"rolled back: {report.reason}" if report else "not applied"
    rolled_back = [(p, why) for p in plan.accepted if all(p is not r for r, _ in plan.rejected)]
    return TailorResult(
        out_dir / f"cv{base.suffix}", out_dir / "cv.pdf", [], plan.rejected + rolled_back, report
    )


def tailor_docx(
    base_docx: Path,
    base_pdf: Path,
    patches: list[Patch],
    rules: Rules,
    out_dir: Path,
    *,
    min_ssim: float,
    render: Renderer = render_pdf,
) -> TailorResult:
    texts = {s.id: s.text for s in load_segments(Document(str(base_docx)))}

    def build(attempt: list[Patch], dest: Path) -> tuple[list[Patch], list[tuple[Patch, str]]]:
        doc = Document(str(base_docx))
        segments = {s.id: s for s in load_segments(doc)}
        applied, failed = [], []
        for p in attempt:
            try:
                apply_patch(segments[p.segment_id], p.old_text, p.new_text)
                applied.append(p)
            except PatchError as e:
                failed.append((p, str(e)))
        if applied:
            doc.save(str(dest))
        return applied, failed

    return _tailor(base_docx, base_pdf, texts, patches, rules, out_dir, build, min_ssim, render)


def tailor_latex(
    base_tex: Path,
    base_pdf: Path,
    patches: list[Patch],
    rules: Rules,
    out_dir: Path,
    *,
    min_ssim: float,
    render: Renderer = render_pdf_latex,
) -> TailorResult:
    """Same as `tailor_docx` for a LaTeX CV; `TailorResult.docx` holds the patched .tex."""
    source = base_tex.read_bytes().decode("utf-8")
    segments = latex.load_segments(source)
    texts = {s.id: s.text for s in segments}

    def build(attempt: list[Patch], dest: Path) -> tuple[list[Patch], list[tuple[Patch, str]]]:
        applied: list[Patch] = []
        failed: list[tuple[Patch, str]] = []
        for p in attempt:  # a patch that would touch anything but text is rejected alone
            try:
                latex.apply_patches(source, segments, [*applied, p])
                applied.append(p)
            except PatchError as e:
                failed.append((p, str(e)))
        if applied:
            dest.write_bytes(latex.apply_patches(source, segments, applied).encode("utf-8"))
        return applied, failed

    return _tailor(base_tex, base_pdf, texts, patches, rules, out_dir, build, min_ssim, render)


async def tailor_for(
    profile: Profile,
    locked_facts: list[str],
    title: str,
    company: str,
    description: str,
    provider: LLMProvider,
    settings: Settings,
    out_dir: Path,
    render: Renderer | None = None,
) -> TailorResult:
    base = Path(profile.cv_path)
    is_tex = base.suffix.lower() == ".tex"
    render = render or renderer_for(base)
    segments = cv_segments(base)
    rules = Rules(
        locked_facts=locked_facts,
        allowed_terms=vocabulary(" ".join(s.text for s in segments), profile.skills),
        max_growth=settings.tailor_max_growth,
        max_word_change=settings.tailor_max_word_change,
    )
    prompt = render_prompt(
        "tailor.j2",
        segments=segments,
        title=title,
        company=company,
        description=description[:8000],
        locked=locked_facts,
        skills=profile.skills,
        max_growth=f"{rules.max_growth:.0%}",
        max_change=f"{rules.max_word_change:.0%}",
    )
    patches = (await provider.complete(prompt, PatchList)).patches
    base_pdf = await asyncio.to_thread(ensure_base_pdf, base, render)
    return await asyncio.to_thread(
        tailor_latex if is_tex else tailor_docx, base, base_pdf, patches, rules, out_dir,
        min_ssim=settings.tailor_min_ssim, render=render,
    )  # fmt: skip


def load_profile(db: sessionmaker[Session]) -> tuple[Profile, list[str]]:
    with db() as s:
        profile = s.get(Profile, 1)
        if profile is None:
            raise TailorError("No profile yet: run `linkit profile import profile.json`")
        if not profile.cv_approved:
            raise TailorError("Approve the converted CV in Setup first")
        return profile, list(s.scalars(select(LockedFact.text)))


async def tailor_job(
    db: sessionmaker[Session],
    job_id: int,
    provider: LLMProvider,
    settings: Settings,
    render: Renderer | None = None,
) -> TailoredCV:
    """Pipeline stage queued_apply -> tailored. Idempotent: an existing tailored CV is reused
    (a job coming back from the Action Center is not tailored, and paid for, twice)."""
    profile, facts = load_profile(db)
    with db.begin() as s:
        job = s.get(Job, job_id)
        if job is None:
            raise TailorError(f"job {job_id} not found")
        latest = s.scalars(
            select(TailoredCV).filter_by(job_id=job_id).order_by(TailoredCV.id.desc())
        ).first()
        if latest is not None and Path(latest.pdf_path).exists():
            transition(s, job, JobState.tailored, "reusing the tailored CV")
            return latest
        title, company, description = job.title, job.company, job.description

    out_dir = settings.data_dir / "tailored" / f"job_{job_id}"
    result = await tailor_for(
        profile, facts, title, company, description, provider, settings, out_dir, render
    )
    with db.begin() as s:
        job = s.get_one(Job, job_id)
        row = TailoredCV(
            job_id=job_id,
            docx_path=str(result.docx),
            pdf_path=str(result.pdf),
            applied=[p.model_dump() for p in result.applied],
            rejected=[{**p.model_dump(), "reason": r} for p, r in result.rejected],
            verify=result.verify.as_dict() if result.verify else {},
        )
        s.add(row)
        transition(
            s,
            job,
            JobState.tailored,
            f"{len(result.applied)} patches applied, {len(result.rejected)} rejected",
            [str(result.docx), str(result.pdf)],
        )
    return row
