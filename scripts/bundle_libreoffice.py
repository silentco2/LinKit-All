"""Build a minimal, portable LibreOffice that only does DOCX -> PDF.

Windows for now: `msiexec /a` extracts the official MSI without installing it, then
whole component groups are removed one at a time. A group stays removed only if every
sample still renders identically to the full LibreOffice (same pages, SSIM >= 0.999).

    uv run python scripts/bundle_libreoffice.py build/lo/dl/LibreOffice_X_Win_x86-64.msi \
        --out build/lo [--sample my_cv.docx ...]

Output: <out>/minimal (the portable tree), <out>/manifest.json and
<out>/libreoffice-minimal-<version>-win-x64.zip.
"""

import argparse
import json
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

import docx
from docx.enum.text import WD_BREAK
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
from docx.shared import Pt
from PIL import Image, ImageDraw

from linkit_all.tailor.verify.layout import verify_pdf

# Removal candidates, tried in order, as globs relative to the LibreOffice root.
# Never candidates: fonts, license files, mergedlo/swlo/writer filters, ICU and locale
# data, CTL/CJK registry (Arabic and Asian CVs must keep rendering right).
P = "program/"
GROUPS: dict[str, list[str]] = {
    "docs-help": ["help", "readmes", "CREDITS.fodt", "update-settings.ini"],
    "presets": ["presets"],
    "dictionaries-non-en": ["share/extensions/dict-[!e]*", "share/extensions/dict-e[!n]*",
                            "share/extensions/nlpsolver", "share/extensions/wiki-publisher"],
    "ui-translations": [P + "resource/*"],
    "ui-langpacks": ["share/registry/Langpack-[!e]*.xcd", "share/registry/Langpack-e[!n]*.xcd",
                     "share/registry/Langpack-en-[!U]*.xcd",
                     "share/registry/res/fcfg_langpack_[!e]*", "share/registry/res/*_[!e]*.xcd"],
    "icon-themes": ["share/config/images_*.zip"],
    "ui-definitions": ["share/config/soffice.cfg", "share/config/wizard"],
    "ui-modules-non-writer": ["share/config/soffice.cfg/modules/[!s]*",
                              "share/config/soffice.cfg/modules/s[!w]*"],
    "gallery-templates": [f"share/{d}" for d in (
        "gallery", "template", "wizards", "autotext", "labels", "tipoftheday", "toolbarmode",
        "fingerprint", "numbertext", "glade", "themes", "theme_definitions", "wordbook",
        "classification", "palette")],
    "basic-scripts": ["share/basic", "share/Scripts", P + "wizards", P + "*.py", P + "*.pyi"],
    "xslt-filters": ["share/xslt", "share/registry/xsltfilter.xcd"],
    "python": [P + "python-core-*", P + "python*.dll", P + "python.exe", P + "pyuno.pyd",
               P + "pythonloader*"],
    "java": [P + "classes", P + "JREProperties.class", P + "java*", P + "jvm*", P + "jdbclo.dll",
             P + "hsqldb.dll", P + "officebean.dll", P + "jpip*", P + "juh*"],
    "dotnet-cli": [P + "cli_*", P + "policy.1.0.*"],
    "calc": [P + "sc*lo.dll", P + "calclo.dll", P + "scalc.exe", P + "CoinMP.dll",
             P + "lpsolve55.dll", P + "orcus*.dll", P + "vbaobjlo.dll", "share/calc",
             "share/registry/calc.xcd"],
    "impress-draw": [P + "sd*lo.dll", P + "simpress.exe", P + "sdraw.exe",
                     P + "PresentationMinimizerlo.dll", "share/registry/impress.xcd",
                     "share/registry/draw.xcd", "share/registry/ogltrans.xcd"],
    "gpu": [P + "opengl", P + "opencl", P + "opencltest.exe", P + "epoxy.dll",
            "share/opengl", "share/skia"],
    "skia": [P + "skialo.dll"],
    "math": [P + "smdlo.dll", P + "smlo.dll", P + "smath.exe", "share/registry/math.xcd"],
    "database": [P + "abplo.dll", P + "dbplo.dll", P + "Engine12.dll", P + "ifbclient.dll",
                 P + "sbase.exe", P + "postgresql-sdbc-impllo.dll", P + "libpq.dll",
                 P + "odbcconfig.exe", "share/firebird", "share/registry/base.xcd",
                 "share/registry/postgresql.xcd", "share/registry/reportbuilder.xcd"],
    "other-import-filters": [P + n for n in (
        "mwaw.dll", "etonyek.dll", "wps.dll", "staroffice.dll", "wpd.dll", "wpg.dll",
        "revenge.dll", "odfgen.dll", "xpdfimport.exe", "pdfimportlo.dll")]
        + ["share/xpdfimport", "share/registry/pdfimport.xcd"],
    "pdfium": [P + "pdfiumlo.dll"],
    "desktop-integration": [P + n for n in (
        "gengal.exe", "gpgme-w32spawn.exe", "mar.exe", "minidump_upload.exe", "quickstart.exe",
        "regview.exe", "senddoc.exe", "spsupp_*", "sweb.exe", "swriter.exe", "twain32shim.exe",
        "uno.exe", "unoinfo.exe", "unopkg*", "updater*", "update_service.exe",
        "soffice_safe.exe", "intro*.png", "shlxthdl", "so_activex.dll", "inprocserv.dll",
        "jumplistlo.dll", "UAccCOM.dll", "winaccessibility.dll", "updchklo.dll",
        "updatecheckuilo.dll", "LanguageToollo.dll", "deploymentgui.dll", "migrationoo*")]
        + ["share/registry/onlineupdate.xcd", "share/registry/librelogo.xcd"],
    "crypto-signing": [P + n for n in (
        "nss3.dll", "nssckbi.dll", "nssdbm3.dll", "nssutil3.dll", "nspr4.dll", "plc4.dll",
        "plds4.dll", "ssl3.dll", "smime3.dll", "softokn3.dll", "freebl3.dll", "libxmlsec*",
        "gpgmepp.dll", "Argon2OptDll.dll")],
    "network": [P + "libcurl.dll", P + "libssl-3.dll", P + "libcrypto-3.dll"],
    "rdf-search": [P + "clucene.dll", P + "librdf.dll", P + "raptor2.dll", P + "rasqal.dll"],
}  # fmt: skip


def prepare(root: Path) -> None:
    """Make the extracted tree self-contained before pruning."""
    # VC++ runtime next to soffice, so users need no redistributable installed.
    for dll in (root / "System64").glob("*.dll"):
        shutil.copy2(dll, root / "program" / dll.name)
    # Bundled fonts (Carlito, Caladea, Liberation, Amiri...) as app-private fonts.
    fonts = root / "share/fonts/truetype"
    fonts.mkdir(parents=True, exist_ok=True)
    for font in (root / "Fonts").glob("*"):
        shutil.move(font, fonts / font.name)
    for leftover in ("System", "System64", "Fonts"):
        shutil.rmtree(root / leftover, ignore_errors=True)
    for msi in root.glob("*.msi"):
        msi.unlink()


NS = (
    'xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006" '
    'xmlns:wps="http://schemas.microsoft.com/office/word/2010/wordprocessingShape" '
    'xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing" '
    'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" '
    'xmlns:v="urn:schemas-microsoft-com:vml"'
)
BOX = "<w:txbxContent><w:p><w:r><w:t>Text box: Python, SQL</w:t></w:r></w:p></w:txbxContent>"
# A Word-style text box: DrawingML shape, plus the VML copy older readers use.
TEXTBOX = (
    f"<w:r {nsdecls('w')} {NS}><mc:AlternateContent>"
    '<mc:Choice Requires="wps"><w:drawing><wp:inline distT="0" distB="0" distL="0" distR="0">'
    '<wp:extent cx="2540000" cy="508000"/><wp:docPr id="10" name="Text Box 10"/>'
    "<a:graphic>"
    '<a:graphicData uri="http://schemas.microsoft.com/office/word/2010/wordprocessingShape">'
    '<wps:wsp><wps:cNvSpPr txBox="1"/><wps:spPr>'
    '<a:xfrm><a:off x="0" y="0"/><a:ext cx="2540000" cy="508000"/></a:xfrm>'
    '<a:prstGeom prst="rect"><a:avLst/></a:prstGeom></wps:spPr>'
    f"<wps:txbx>{BOX}</wps:txbx><wps:bodyPr/></wps:wsp>"
    "</a:graphicData></a:graphic></wp:inline></w:drawing></mc:Choice>"
    '<mc:Fallback><w:pict><v:shape style="width:200pt;height:40pt" type="#_x0000_t202">'
    f"<v:textbox>{BOX}</v:textbox></v:shape></w:pict></mc:Fallback>"
    "</mc:AlternateContent></w:r>"
)


def make_sample(path: Path) -> Path:
    """A CV-like DOCX touching what CVs use: fonts, header/footer, table, list, image, box."""
    img = path.with_suffix(".png")
    im = Image.new("RGB", (120, 120), "white")
    ImageDraw.Draw(im).ellipse((10, 10, 110, 110), fill="steelblue")
    im.save(img)

    d = docx.Document()
    d.sections[0].header.paragraphs[0].text = "Jane Doe · jane@example.com · +1 555 0100"
    d.sections[0].footer.paragraphs[0].text = "References available on request"
    d.add_picture(str(img))
    for font in ("Calibri", "Arial", "Times New Roman", "Cambria"):
        r = d.add_paragraph().add_run(f"{font}: Senior Software Engineer, Acme Corp, 2019-2023")
        r.font.name, r.font.size = font, Pt(11)
    p = d.add_paragraph()
    p.add_run("Bold ").bold = True
    p.add_run("italic ").italic = True
    p.add_run("underlined").underline = True
    for item in ("Built scalable APIs", "Led a team of 5", "Cut latency by 40%"):
        d.add_paragraph(item, style="List Bullet")
    t = d.add_table(rows=2, cols=2, style="Table Grid")
    for i, txt in enumerate(("Skill", "Years", "Python", "6")):
        t.cell(i // 2, i % 2).text = txt
    d.add_paragraph()._p.append(parse_xml(TEXTBOX))
    d.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    d.add_paragraph("Page two: education, B.Sc. Computer Science")
    d.save(str(path))
    return path


def convert(soffice: Path, src: Path, out_dir: Path) -> Path | None:
    profile = Path(tempfile.mkdtemp(prefix="lo-profile-"))
    try:
        subprocess.run(  # noqa: S603 (soffice we just extracted)
            [str(soffice), f"-env:UserInstallation={profile.as_uri()}", "--headless",
             "--norestore", "--convert-to", "pdf", "--outdir", str(out_dir), str(src)],
            capture_output=True, timeout=180, check=False,
        )  # fmt: skip
    except subprocess.TimeoutExpired:
        return None
    finally:
        shutil.rmtree(profile, ignore_errors=True)
    pdf = out_dir / f"{src.stem}.pdf"
    return pdf if pdf.exists() else None


def renders_same(soffice: Path, samples: list[Path], refs: dict[Path, Path], work: Path) -> bool:
    out = work / "check"
    shutil.rmtree(out, ignore_errors=True)
    out.mkdir(parents=True)
    for s in samples:
        pdf = convert(soffice, s, out)
        if pdf is None or not verify_pdf(refs[s], pdf, min_ssim=0.999).ok:
            return False
    return True


def size_mb(root: Path) -> float:
    return sum(f.stat().st_size for f in root.rglob("*") if f.is_file()) / 1e6


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("msi", type=Path)
    ap.add_argument("--out", type=Path, default=Path("build/lo"))
    ap.add_argument("--sample", type=Path, action="append", default=[])
    args = ap.parse_args()
    out: Path = args.out.resolve()
    version = args.msi.stem.split("_")[1]

    full = out / "full"  # untouched admin image, the reference
    if not (full / "program/soffice.exe").exists():
        print("Extracting MSI (no install)...")
        msiexec = shutil.which("msiexec") or sys.exit("msiexec not found (Windows only)")
        subprocess.run(  # noqa: S603
            [msiexec, "/a", str(args.msi.resolve()), "/qn", f"TARGETDIR={full}"], check=True
        )

    work = out / "work"
    shutil.rmtree(work, ignore_errors=True)
    (work / "ref").mkdir(parents=True)
    samples = [make_sample(work / "sample.docx"), *(s.resolve() for s in args.sample)]
    refs = {}
    for s in samples:
        ref = convert(full / "program/soffice.exe", s, work / "ref")
        if ref is None:
            sys.exit(f"Full LibreOffice failed to convert {s}")
        refs[s] = ref

    print("Copying and preparing the minimal tree...")
    minimal = out / "minimal"
    shutil.rmtree(minimal, ignore_errors=True)
    shutil.copytree(full, minimal)
    prepare(minimal)
    soffice = minimal / "program/soffice.exe"
    if not renders_same(soffice, samples, refs, work):
        sys.exit("Prepared tree already renders differently from the full one; aborting")
    trash = work / "trash"
    removed: list[str] = []
    kept: list[str] = []

    def try_remove(label: str, paths: list[Path]) -> bool:
        moved = []
        for f in sorted(paths, reverse=True):
            if not f.exists():
                continue
            dest = trash / label.replace(":", "_").replace("/", "_") / f.relative_to(minimal)
            dest.parent.mkdir(parents=True, exist_ok=True)
            shutil.move(f, dest)
            moved.append((f, dest))
        if not moved:
            return False
        if renders_same(soffice, samples, refs, work):
            removed.append(label)
            print(f"  removed {label} ({len(moved)} paths)")
            return True
        for f, dest in reversed(moved):
            f.parent.mkdir(parents=True, exist_ok=True)
            shutil.move(dest, f)
        kept.append(label)
        print(f"  KEPT    {label} (needed for rendering)")
        return False

    for name, globs in GROUPS.items():
        try_remove(name, [f for g in globs for f in minimal.glob(g)])

    # Second pass: a kept group may need only some of its files; try them one by one.
    for name in [k for k in kept if ":" not in k]:
        for f in sorted({f for g in GROUPS[name] for f in minimal.glob(g)}):
            try_remove(f"{name}:{f.relative_to(minimal).as_posix()}", [f])
    shutil.rmtree(trash, ignore_errors=True)

    manifest = {
        "version": version,
        "removed_groups": removed,
        "kept_groups": kept,
        "full_mb": round(size_mb(full), 1),
        "minimal_mb": round(size_mb(minimal), 1),
    }
    (out / "manifest.json").write_text(json.dumps(manifest, indent=2))
    archive = shutil.make_archive(
        str(out / f"libreoffice-minimal-{version}-win-x64"), "zip", minimal
    )
    print(json.dumps(manifest, indent=2), "\n", archive)


if __name__ == "__main__":
    main()
